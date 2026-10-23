"""Style the hidden-card / mobile design DOCX."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PATH = Path("docs/product/隐藏牌与手机端设计.docx")
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WIDTHS = {
    0: [3.2, 5.2, 4.8, 3.5],
    1: [3.2, 5.2, 4.3, 4.0],
    2: [3.4, 5.0, 5.2, 3.1],
    3: [8.2, 8.2],
}


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


doc = Document(PATH)
for idx, table in enumerate(doc.tables):
    table.autofit = False
    table.allow_autofit = False
    widths = WIDTHS.get(idx)
    if widths and len(widths) == len(table.columns):
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    if table.rows:
        hdr = table.rows[0]
        hdr._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell in hdr.cells:
            shade(cell, "123B5D")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.color.rgb = WHITE
                    run.font.name = "微软雅黑"
                    run.font.size = Pt(8.2)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    if not run.font.size:
                        run.font.size = Pt(8.2)
                    run.font.name = "宋体"
                    rpr = run._element.get_or_add_rPr()
                    rfonts = rpr.find(qn("w:rFonts"))
                    if rfonts is None:
                        rfonts = OxmlElement("w:rFonts")
                        rpr.append(rfonts)
                    rfonts.set(qn("w:eastAsia"), "宋体")

for shape in doc.inline_shapes:
    try:
        ratio = shape.height / shape.width
    except (ZeroDivisionError, TypeError):
        continue
    shape.width = Cm(9.5)
    shape.height = int(Cm(9.5) * ratio)

doc.core_properties.title = "法前哨 · 隐藏牌与手机端设计"
doc.save(PATH)
print("styled", PATH)
