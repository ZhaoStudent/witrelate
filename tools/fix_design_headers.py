"""Set clean headers/footers on the two product-design DOCX files."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FILES = {
    Path("docs/product/产品形态与界面设计.docx"): ("法前哨 · 产品形态与界面设计", "2026-10-22"),
    Path("docs/product/隐藏牌与手机端设计.docx"): ("法前哨 · 隐藏牌与手机端设计 v2", "2026-10-22"),
}
GRAY = RGBColor(0x8A, 0x99, 0xA8)


def clear_part(part):
    for p in list(part.paragraphs):
        p._element.getparent().remove(p._element)


def style_run(run, text, size=8):
    run.text = text
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    run.font.color.rgb = GRAY
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "微软雅黑")


def add_page_field(paragraph):
    run = paragraph.add_run()
    run.font.name = "微软雅黑"
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "微软雅黑")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


for path, (header_text, date_text) in FILES.items():
    doc = Document(path)
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for part, is_header in ((section.header, True), (section.footer, False)):
            part.is_linked_to_previous = False
            clear_part(part)
            p = part.add_paragraph()
            if is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                style_run(p.add_run(), header_text, 8)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                style_run(p.add_run(), f"{date_text} · 第 ", 8)
                add_page_field(p)
                style_run(p.add_run(), " 页", 8)
    doc.save(path)
    print("headers", path)
