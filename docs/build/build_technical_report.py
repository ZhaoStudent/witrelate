"""Build the styled reference DOCX and the final 技术方案 DOCX.

Run with the bundled Python from `load_workspace_dependencies`.
This script is a build tool, not part of the runtime package.
"""

from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
AICOMP = ROOT / "docs" / "aicomp"
SRC_MD = AICOMP / "技术方案-法前哨.md"
REF_RAW = AICOMP / "_reference.docx"
REF_STYLED = AICOMP / "_reference_styled.docx"
OUT_DOCX = AICOMP / "技术方案-法前哨.docx"

NAVY = RGBColor(0x12, 0x3B, 0x5D)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MID_BLUE = RGBColor(0x2B, 0x5C, 0x8A)
GRAY = RGBColor(0x5B, 0x6B, 0x7C)
LIGHT_GRAY = RGBColor(0x7B, 0x8B, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_style_font(style, name: str, size: float, bold: bool | None = None,
                   color: RGBColor | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_bottom_border(paragraph, color: str = "123B5D", size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def style_reference() -> None:
    if not REF_RAW.exists():
        with REF_RAW.open("wb") as fh:
            subprocess.run(
                ["pandoc", "--print-default-data-file", "reference.docx"],
                stdout=fh,
                check=True,
            )
    doc = Document(REF_RAW)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(0.9)

    body = ["Body Text", "First Paragraph", "Compact"]
    for name in body:
        style = doc.styles[name]
        set_style_font(style, "宋体", 9.8)
        pf = style.paragraph_format
        pf.line_spacing = 1.25
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.first_line_indent = Cm(0.7)

    normal = doc.styles["Normal"]
    set_style_font(normal, "宋体", 9.8)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(2)

    title = doc.styles["Title"]
    set_style_font(title, "微软雅黑", 27, bold=True, color=NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(42)
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, "微软雅黑", 12.5, bold=False, color=GRAY)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)

    for name, size, color, before, after in [
        ("Heading 1", 16, NAVY, 11, 5),
        ("Heading 2", 13, BLUE, 9, 3),
        ("Heading 3", 11.5, MID_BLUE, 7, 2),
        ("Heading 4", 10.5, MID_BLUE, 5, 2),
    ]:
        style = doc.styles[name]
        set_style_font(style, "微软雅黑", size, bold=True, color=color)
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.line_spacing = 1.1

    for name in ["Caption", "Image Caption", "Table Caption"]:
        style = doc.styles[name]
        set_style_font(style, "微软雅黑", 8.2, bold=False, color=LIGHT_GRAY)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(1)
        style.paragraph_format.space_after = Pt(4)

    block = doc.styles["Block Text"]
    set_style_font(block, "宋体", 9.5, color=GRAY)
    block.paragraph_format.left_indent = Cm(0.5)
    block.paragraph_format.space_after = Pt(2)

    for name in ["Source Code", "Verbatim Char"]:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        set_style_font(style, "Consolas", 8.2)
        style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.save(REF_STYLED)


def convert() -> None:
    cmd = [
        "pandoc",
        str(SRC_MD),
        "-o",
        str(OUT_DOCX),
        f"--reference-doc={REF_STYLED}",
        f"--resource-path={AICOMP}",
        "--from=markdown+pipe_tables+task_lists",
        "--columns=100",
    ]
    subprocess.run(cmd, check=True, cwd=ROOT)


def postprocess() -> None:
    doc = Document(OUT_DOCX)

    # Cover: remove the Markdown title block, then insert a cover page.
    body_start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("作品简介"):
            body_start = i
            break
    if body_start is None:
        raise RuntimeError("找不到正文起点：作品简介")
    for p in list(doc.paragraphs[:body_start]):
        p._element.getparent().remove(p._element)

    first = doc.paragraphs[0]
    cover_items = [
        ("法前哨", "Title"),
        ("面向法律纠纷“事前准备”的可审计多智能体系统", "Subtitle"),
        ("技术方案", "Subtitle"),
    ]
    for text, style_name in cover_items:
        p = first.insert_paragraph_before(text, style=style_name)
        if text == "技术方案":
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(26)

    meta = first.insert_paragraph_before(
        "参赛主题：算法主题赛 · 智能体互联    作品名称：法前哨\n"
        "平台：梧桐智能体互联平台（wt.ioa.pub）    版本：v1.0 · 2026-10-21",
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(10)
    for run in meta.runs:
        run.font.name = "微软雅黑"
        run.font.size = Pt(9.5)
        run.font.color.rgb = GRAY
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    rule = first.insert_paragraph_before("")
    add_bottom_border(rule, "123B5D", 10)
    rule.paragraph_format.space_after = Pt(16)

    outline = first.insert_paragraph_before(
        "一、作品概述    二、需求分析    三、技术方案\n"
        "四、项目实施    五、应用成效    六、总结与展望    附录",
    )
    outline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    outline.paragraph_format.line_spacing = 1.8
    outline.paragraph_format.space_after = Pt(18)
    for run in outline.runs:
        run.font.name = "微软雅黑"
        run.font.size = Pt(10)
        run.font.color.rgb = MID_BLUE
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    page_break = first.insert_paragraph_before("")
    page_break.add_run().add_break(WD_BREAK.PAGE)

    # Tables: header shading, repeated header, sensible column widths.
    table_widths = {
        0: [2.6, 4.3, 5.6, 4.0],       # 1.2 四个能力
        1: [3.5, 4.8, 8.2],            # 2.1 三类用户
        2: [0.8, 3.0, 5.7, 7.0],       # 2.3 五个缺口
        3: [2.5, 4.3, 9.7],            # 2.4 需求指标
        4: [3.8, 2.0, 6.2, 4.5],       # 3.2 角色与契约
        5: [8.25, 8.25],               # 3.4 复用与自创
        6: [3.8, 6.4, 6.3],            # 3.5 兼容性
        7: [2.8, 11.7, 2.0],           # 5.3 必要验证
        8: [4.4, 2.2, 2.2, 2.2, 5.5],  # 5.4 量化指标
        9: [3.8, 5.0, 7.5],            # 6.2 不足
        10: [6.0, 10.5],               # E. 关键代码入口
    }
    for t_idx, table in enumerate(doc.tables):
        table.autofit = False
        table.allow_autofit = False
        widths = table_widths.get(t_idx)
        if widths and len(widths) == len(table.columns):
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        if table.rows:
            hdr = table.rows[0]
            tr_pr = hdr._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
            for cell in hdr.cells:
                shade(cell, "123B5D")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                        run.font.size = Pt(8.2)
                        run.font.name = "微软雅黑"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    for run in p.runs:
                        if not run.font.size:
                            run.font.size = Pt(8.2)
                        if not run.font.name:
                            run.font.name = "宋体"
                        rpr = run._element.get_or_add_rPr()
                        rfonts = rpr.find(qn("w:rFonts"))
                        if rfonts is None:
                            rfonts = OxmlElement("w:rFonts")
                            rpr.append(rfonts)
                        rfonts.set(qn("w:eastAsia"), "宋体")

    # Images: varied widths so screenshots stay legible but don't eat pages.
    widths = [Cm(15.5), Cm(12.5), Cm(11.0), Cm(16.0)]
    for idx, shape in enumerate(doc.inline_shapes):
        try:
            ratio = shape.height / shape.width
        except (ZeroDivisionError, TypeError):
            continue
        width = widths[idx] if idx < len(widths) else Cm(11.0)
        shape.width = width
        shape.height = int(width * ratio)

    # Keep each image with its caption.
    for p in doc.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
        if p.text.strip() == "附录":
            p.paragraph_format.page_break_before = True

    # Header / footer; no header on cover.
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("法前哨 · 技术方案 v1.0")
    hr.font.name = "微软雅黑"
    hr.font.size = Pt(8)
    hr.font.color.rgb = LIGHT_GRAY
    hr._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("2026-10-21 · 第 ")
    add_page_field(fp)
    r3 = fp.add_run(" 页")
    for run in (r1, r3):
        run.font.name = "微软雅黑"
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT_GRAY
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    style_reference()
    convert()
    postprocess()
    print("built", OUT_DOCX)
